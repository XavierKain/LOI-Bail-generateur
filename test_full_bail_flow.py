"""
Test complet du flux de génération BAIL avec formatage des placeholders.
"""

from docx import Document
from modules.word_text_loader import WordTextLoader
from modules.placeholder_formatter import replace_placeholder_with_format
import re

print('=' * 80)
print('TEST COMPLET DU FLUX BAIL AVEC FORMATAGE')
print('=' * 80)

# 1. Charger les textes formatés depuis Word
print('\n1️⃣  Chargement des textes formatés depuis Word...')
loader = WordTextLoader("Textes BAIL avec styles.docx")
print(f'   ✅ {len(loader.get_section_ids())} sections chargées')

# 2. Créer un document de test
print('\n2️⃣  Création du document de test...')
doc = Document()

# 3. Simuler l'ajout d'une section avec placeholder en gras
print('\n3️⃣  Test avec section COMPARUTION_COMPARUTION_PRENEUR (placeholders en gras)...')

# Récupérer une section de preneur (qui devrait avoir des placeholders en gras)
section_id = "COMPARUTION_COMPARUTION_PRENEUR"
source_para = loader.get_formatted_paragraph(section_id)

if source_para:
    print(f'   📄 Texte source: {source_para.text[:100]}...')
    print(f'   📊 Runs dans le source: {len(source_para.runs)}')

    # Analyser le formatage des placeholders
    print('\n   🔍 Analyse du formatage des placeholders:')
    for i, run in enumerate(source_para.runs):
        if '[' in run.text or ']' in run.text:
            print(f'      Run {i}: "{run.text}" (bold={run.bold})')

    # Copier dans le document de test
    target_para = doc.add_paragraph()
    loader.copy_formatted_text_to_paragraph(source_para, target_para)

    print(f'\n   ✅ Texte copié dans le document')
    print(f'   📊 Runs dans la cible: {len(target_para.runs)}')

    # Vérifier que les placeholders sont bien en gras
    print('\n   🔍 Vérification du formatage après copie:')
    full_text = target_para.text
    placeholders = re.findall(r'\[([^\]]+)\]', full_text)
    print(f'      Placeholders trouvés: {placeholders}')

    for i, run in enumerate(target_para.runs):
        if '[' in run.text or ']' in run.text:
            print(f'      Run {i}: "{run.text}" (bold={run.bold})')

    # 4. Simuler le remplacement des placeholders
    print('\n4️⃣  Remplacement des placeholders...')

    # Données de test
    donnees = {
        'Dénomination du preneur': 'SARL DUPONT AUTOMOBILES',
        'Forme juridique du preneur': 'SARL',
        'Capital du preneur': '50000',
        'Adresse du siège social du preneur': '123 Avenue de la République',
        'Code postal du siège social du preneur': '75011',
        'Ville du siège social du preneur': 'PARIS',
        'SIREN du preneur': '123 456 789',
        'RCS du preneur': 'Paris'
    }

    # Remplacer chaque placeholder
    for placeholder_match in placeholders:
        placeholder_with_brackets = f"[{placeholder_match}]"

        # Trouver la valeur correspondante
        value = donnees.get(placeholder_match, f"MANQUANT_{placeholder_match}")

        print(f'      Remplace "{placeholder_with_brackets}" par "{value}"')
        success = replace_placeholder_with_format(target_para, placeholder_with_brackets, value)
        print(f'         → {"✅ OK" if success else "❌ ÉCHEC"}')

    # 5. Vérifier le résultat final
    print('\n5️⃣  Résultat final:')
    print(f'   📄 Texte: {target_para.text[:150]}...')
    print(f'   📊 Runs: {len(target_para.runs)}')

    print('\n   🔍 Formatage des runs finaux:')
    for i, run in enumerate(target_para.runs):
        if run.text.strip():  # Ne montrer que les runs non vides
            preview = run.text[:50] if len(run.text) > 50 else run.text
            print(f'      Run {i}: "{preview}" (bold={run.bold})')

    # Sauvegarder
    doc.save('test_full_bail_flow_output.docx')
    print('\n💾 Document sauvegardé: test_full_bail_flow_output.docx')

else:
    print(f'   ❌ Section {section_id} non trouvée')

print('\n' + '=' * 80)
print('TEST TERMINÉ')
print('=' * 80)
